from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "format.docx"
OUTPUT = ROOT / "docs" / "AI_Interview_Platform_Project_Analysis.docx"


PARAGRAPHS = [
    "AI Interview Platform",
    "",
    "1. Project Overview & Scope",
    "1.1 Problem Statement",
    "Recruiters and interviewers often spend substantial time creating role-specific questions, coordinating candidate interviews, reviewing recordings, and comparing responses. Manual workflows make it difficult to run consistent first-round interviews at scale and provide timely, evidence-based feedback.",
    "The AI Interview Platform addresses this problem with a browser-based interviewer portal and candidate interview flow. Interviewers can create sessions from role information, generate structured questions, share a unique link, and review candidate responses. Candidates can register, provide consent, record video or audio answers, and receive speech-to-text processing where the browser supports it.",
    "1.2 Objectives",
    "The primary objectives of the AI Interview Platform are to:",
    "1.2.1 Enable Interview Session Management: Provide interviewer screens for creating, editing, publishing, and deleting interview sessions with configurable duration, questions, media options, and share links.",
    "1.2.2 Generate Role-Specific Questions: Use the server-side AI integration to create six to ten structured questions from a job description, role title, and optional candidate resume details.",
    "1.2.3 Support Candidate Onboarding: Allow candidates to open a unique interview link, register their profile, optionally submit resume information, review requirements, and provide recording consent.",
    "1.2.4 Conduct Recorded Interviews: Present questions through the candidate interview room while capturing video and audio through browser media APIs and tracking question timing and progress.",
    "1.2.5 Provide Speech and AI Evaluation: Transcribe spoken answers when Web Speech API support is available and analyze responses against expected keywords and evaluation criteria.",
    "1.2.6 Monitor Interview Integrity: Collect probabilistic phone, gaze, audio, server, and tamper-related signals and expose evidence for human review rather than automatic rejection.",
    "1.2.7 Persist and Synchronize Results: Support browser localStorage for standalone development and synchronize sessions, candidates, and integrity events with the optional Node.js API, JSON storage, or MongoDB.",
    "1.2.8 Provide Operational Deployment Paths: Support local Vite development, LAN testing with automatic base URLs, a lightweight Node.js backend, and Vercel serverless API deployment.",
    "",
    "",
    "1.3 Scope",
    "Users of the System:",
    "Interviewer / HR Administrator: Creates sessions, configures questions and media, trains the avatar profile, shares links, and reviews candidate scores and rankings.",
    "Candidate: Opens a shared interview link, registers, grants permissions and consent, answers questions, and completes the interview.",
    "AI and Integrity Services: Generate questions, analyze resumes and responses, and receive integrity event batches for interviewer review.",
    "Current Scope (Phase 1):",
    "React 18 single-page application with React Router routes for interviewer and candidate workflows.",
    "Interviewer dashboard, sessions list, session creation and detail screens, candidate list, avatar training, and settings.",
    "Candidate registration, interview room, completion screen, and phone recording route.",
    "Question generation from job descriptions and resume context through a server-side OpenAI-compatible JSON API.",
    "Video and audio recording through MediaRecorder, browser speech recognition, timer and progress controls.",
    "Response, resume, and integrity analysis endpoints with structured JSON validation and error handling.",
    "Local browser persistence with optional synchronization to the Node.js backend, JSON file storage, or MongoDB.",
    "Future Scope (Later Phases):",
    "Replace the local no-auth interviewer profile with production identity, role-based access control, and organization-level tenancy.",
    "Add production media object storage, retention policies, encrypted recording delivery, and candidate notification workflows.",
    "Expand automated evaluation calibration, reporting, accessibility coverage, and analytics across multiple interview campaigns.",
    "1.4 Development Methodology / Process Planning",
    "The platform follows an iterative Agile development approach. Features are organized around the interviewer and candidate workflows, with reusable React components and service modules supporting incremental delivery. Local-first behavior enables rapid UI development, while optional backend persistence and AI services are integrated as deployment and evaluation needs grow.",
    "",
    "2. Phase 2 – Requirement Analysis",
    "Functional Requirements",
    "Module I: Interviewer Session & Avatar Management",
    "FR-01 (Session Creation): The system shall allow an interviewer to create a session with role information, questions, duration, and interview settings.",
    "FR-02 (Question Configuration): The system shall allow an interviewer to add, edit, delete, reorder, and configure questions with type, time limit, weight, keywords, and evaluation criteria.",
    "FR-03 (AI Question Generation): The system shall generate structured questions from job descriptions, role context, and optional resume data through a server-side AI endpoint.",
    "FR-04 (Share Link): The system shall generate and display a unique candidate interview link for a saved session.",
    "Module II: Candidate Registration & Interview Room",
    "FR-05 (Candidate Registration): The system shall register a candidate against the session identified by the shared link.",
    "FR-06 (Consent and Permissions): The system shall request candidate consent and browser permission before recording video or audio.",
    "FR-07 (Interview Capture): The system shall present questions, track progress and time limits, and record candidate video or audio according to session settings.",
    "FR-08 (Speech Transcription): The system shall transcribe candidate speech when Web Speech API support is available and retain the answer text with the response.",
    "Module III: Evaluation, Results & Integrity",
    "FR-09 (Resume Analysis): The system shall accept resume text or profile data and make structured analysis available to the interview workflow.",
    "FR-10 (Answer Analysis): The system shall evaluate responses against question context, expected keywords, and evaluation criteria and return structured scores and feedback.",
    "FR-11 (Integrity Events): The system shall collect client and server integrity events, calculate a weighted evidence score, and record that signals are probabilistic.",
    "FR-12 (Candidate Results): The system shall display candidate scores, rankings, recordings, response analysis, and integrity evidence to the interviewer.",
    "Module IV: Persistence, API & Deployment",
    "FR-13 (Local Persistence): The system shall support localStorage persistence for interviewer profiles, sessions, candidates, and development fallback behavior.",
    "FR-14 (Remote Persistence): The system shall provide session and candidate APIs backed by MongoDB when configured and JSON file storage otherwise.",
    "FR-15 (Service Health): The backend shall expose a health endpoint that reports availability and the active persistence mode.",
    "FR-16 (Deployment Configuration): The system shall support Vite development, LAN base URL configuration, and Vercel rewrites/serverless API deployment settings.",
    "",
    "3. Phase 3 – System Design & Modeling",
    "3.1 Use Case Diagram",
    "",
    "3.2 Activity Diagram",
    "",
    "",
    "",
    "",
]


TABLE_ROWS = [
    ["Iteration", "Sprint Focus", "Delivered Features / Architectural Modules"],
    ["Iteration 1", "Core React UI & Routing", "Vite application shell, React Router routes, interviewer and candidate layouts, shared controls, error boundary."],
    ["Iteration 2", "Session & Candidate Workflows", "Session creation and editing, unique links, candidate registration, interviewer dashboard, session and candidate views."],
    ["Iteration 3", "Media & Interview Room", "Camera and microphone access, MediaRecorder capture, timers, speech-to-text integration, progress and completion flow."],
    ["Iteration 4", "AI Question & Response Services", "Structured question generation, resume analysis, response scoring, OpenAI-compatible server handlers, validation and error responses."],
    ["Iteration 5", "Integrity & Review", "Phone, gaze, audio, server and tamper signals, weighted integrity scoring, event upload, candidate ranking and review screens."],
    ["Iteration 6", "Persistence & Multi-Device Support", "localStorage fallback, Node.js HTTP backend, JSON store, MongoDB adapter, synchronization and session-by-link lookup."],
    ["Iteration 7", "Integration & Deployment", "LAN launch scripts, configurable API and public URLs, Vercel configuration, health endpoint, production build and lint scripts."],
]


def replace_template_images(document):
    image_paths = [
        ROOT / "diagrams" / "diagram_usecase_perfect.png",
        ROOT / "diagrams" / "diagram_activity_perfect.png",
    ]
    image_parts = [
        relationship.target_part
        for relationship in document.part.rels.values()
        if relationship.reltype == RT.IMAGE
    ]
    if len(image_parts) != len(image_paths):
        raise RuntimeError(f"Template image count changed: {len(image_parts)}")
    for image_part, image_path in zip(image_parts, image_paths):
        image_part._blob = image_path.read_bytes()


def replace_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(TEMPLATE)
    if len(document.paragraphs) != len(PARAGRAPHS):
        raise RuntimeError(f"Template paragraph count changed: {len(document.paragraphs)}")
    replace_template_images(document)

    for paragraph, text in zip(document.paragraphs, PARAGRAPHS):
        replace_paragraph(paragraph, text)

    table = document.tables[0]
    if len(table.rows) != len(TABLE_ROWS):
        raise RuntimeError(f"Template table row count changed: {len(table.rows)}")
    for row, values in zip(table.rows, TABLE_ROWS):
        for cell, value in zip(row.cells, values):
            replace_paragraph(cell.paragraphs[0], value)

    document.core_properties.title = "AI Interview Platform Project Analysis"
    document.core_properties.subject = "Project analysis and requirements"
    document.core_properties.author = "Project Team"
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()