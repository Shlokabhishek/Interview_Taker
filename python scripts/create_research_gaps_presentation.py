from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


TITLE = "AI Interview Platform"
SUBTITLE = "Title, Abstract, Introduction, Objectives, Existing System, and Research Gaps"
ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "presentations" / "AI_Interview_Platform_Research_Gaps_Presentation.pptx"


NAVY = RGBColor(15, 23, 42)
INDIGO = RGBColor(79, 70, 229)
BLUE = RGBColor(37, 99, 235)
SLATE = RGBColor(51, 65, 85)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 252)
CARD = RGBColor(255, 255, 255)
BORDER = RGBColor(226, 232, 240)
TINT = RGBColor(238, 242, 255)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = INDIGO
    bar.line.fill.background()


def add_title_block(slide, section, title):
    tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(4.5), Inches(0.35))
    p = tag.text_frame.paragraphs[0]
    p.text = section.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = INDIGO

    tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.7))
    p2 = tbox.text_frame.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = NAVY


def add_card(slide, left, top, width, height, title, bullets, tint=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = TINT if tint else CARD
    shape.line.color.rgb = INDIGO if tint else BORDER

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY if tint else INDIGO
    p.space_after = Pt(10)

    for bullet in bullets:
        bp = tf.add_paragraph()
        bp.text = bullet
        bp.level = 0
        bp.bullet = True
        bp.font.size = Pt(13)
        bp.font.color.rgb = SLATE
        bp.space_after = Pt(6)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    add_top_bar(slide)

    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(11.0), Inches(1.4))
    p = title_box.text_frame.paragraphs[0]
    p.text = TITLE
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = title_box.text_frame.add_paragraph()
    p2.text = "An AI-Driven Automated Interview Platform with Avatar-Based Candidate Assessment"
    p2.font.size = Pt(19)
    p2.font.color.rgb = RGBColor(199, 210, 254)
    p2.space_before = Pt(10)

    sub = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.0), Inches(11.2), Inches(1.8)
    )
    sub.fill.solid()
    sub.fill.fore_color.rgb = RGBColor(30, 41, 59)
    sub.line.color.rgb = INDIGO
    tf = sub.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.25)

    a = tf.paragraphs[0]
    a.text = "Research-Focused Presentation"
    a.font.size = Pt(18)
    a.font.bold = True
    a.font.color.rgb = WHITE

    b = tf.add_paragraph()
    b.text = SUBTITLE
    b.font.size = Pt(12)
    b.font.color.rgb = RGBColor(191, 219, 254)
    b.space_before = Pt(8)

    c = tf.add_paragraph()
    c.text = "Prepared from the existing project report and research paper."
    c.font.size = Pt(11)
    c.font.color.rgb = RGBColor(148, 163, 184)
    c.space_before = Pt(8)


def add_content_slide(prs, section, title, left_title, left_bullets, right_title=None, right_bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG)
    add_title_block(slide, section, title)

    if right_title and right_bullets:
        add_card(slide, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), left_title, left_bullets)
        add_card(slide, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), right_title, right_bullets, tint=True)
    else:
        add_card(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), left_title, left_bullets, tint=True)


def add_abstract_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG)
    add_title_block(slide, "Section 1", "Abstract")

    body = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2)
    )
    body.fill.solid()
    body.fill.fore_color.rgb = TINT
    body.line.color.rgb = INDIGO

    tf = body.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    tf.margin_bottom = Inches(0.25)

    abstract_text = (
        "The rapid growth of remote hiring has created demand for interview systems that scale "
        "without sacrificing consistency or candidate experience. This paper presents AI Interview "
        "Platform, a web-based system that lets recruiting teams run automated, asynchronous video "
        "interviews through a personalized AI avatar. The platform combines a React single-page "
        "application with a serverless backend, and uses a large language model (LLM) for question "
        "generation, response evaluation, and candidate ranking, with a deterministic rule-based engine "
        "as a fallback whenever the LLM is unavailable. A distinguishing feature is the ability to "
        "synthesize an interviewer avatar from an HR professional's photographs and voice, paired with "
        "text-to-speech (TTS) and speech-to-text (STT) for a conversational interview flow. The system "
        "also includes a lightweight, privacy-preserving computer-vision heuristic that flags likely phone "
        "use during recording. We describe the architecture, a five-dimension scoring model (relevance, "
        "accuracy, confidence, keyword coverage, and skill signal), and the graceful-degradation strategy "
        "that keeps the platform operational under API outages or rate limits. We report a functional and "
        "qualitative evaluation of the implemented system and, because automated hiring tools now fall "
        "under active legal and regulatory scrutiny (e.g., NYC Local Law 144, the Illinois AI Video "
        "Interview Act, and the EU AI Act's high-risk classification of employment AI), we discuss the "
        "fairness, transparency, and auditability requirements the platform would need to satisfy before "
        "production deployment. We position this as a systems and design contribution rather than a "
        "validated hiring instrument, and we outline the empirical study needed to establish predictive "
        "validity and demographic fairness prior to real-world use."
    )

    p1 = tf.paragraphs[0]
    p1.text = abstract_text
    p1.font.size = Pt(13)
    p1.font.color.rgb = SLATE
    p1.alignment = PP_ALIGN.JUSTIFY
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = (
        "Index Terms: Artificial Intelligence, Large Language Models, Automated Interviewing, "
        "AI Avatar, Speech Processing, Candidate Assessment, Algorithmic Fairness, "
        "Human-Computer Interaction"
    )
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = NAVY


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_abstract_slide(prs)

    add_content_slide(
        prs,
        "Section 2",
        "Introduction",
        "Problem Context",
        [
            "Traditional interviews are time-consuming, difficult to scale, and prone to interviewer bias.",
            "Remote hiring increases candidate volume and makes scheduling and evaluation consistency harder.",
            "Recent NLP and LLM advances enable automatic question generation and response analysis.",
        ],
        "Why This Work Matters",
        [
            "Many current automated tools feel impersonal because they rely on forms or chat-style interfaces.",
            "That weakens soft-skill assessment and reduces the realism of the interview experience.",
            "This work aims to make AI interviewing more human-like while keeping the process scalable and objective.",
        ],
    )

    add_content_slide(
        prs,
        "Section 3",
        "Objectives",
        "Primary Objectives",
        [
            "Develop a full-stack automated interview platform for interviewer and candidate workflows.",
            "Create a personalized AI avatar to deliver questions in a more natural interview format.",
            "Generate role-specific and resume-aware interview questions automatically.",
            "Score responses across relevance, accuracy, confidence, keyword coverage, and professional skills.",
        ],
        "Operational Objectives",
        [
            "Record video, audio, and transcripts for review and ranking.",
            "Detect possible phone usage during interviews using lightweight client-side vision heuristics.",
            "Maintain service continuity with a rule-based fallback when external AI APIs are unavailable.",
        ],
    )

    add_content_slide(
        prs,
        "Section 4",
        "Existing System",
        "Current Approaches",
        [
            "Manual interviews depend heavily on human availability, scheduling, and subjective judgment.",
            "Structured interview systems improve consistency but still require significant interviewer effort.",
            "Earlier automated tools often used fixed question banks and keyword matching.",
            "Modern AI interview tools may analyze answers semantically, but many still provide limited interaction quality.",
        ],
        "Limitations of Existing Systems",
        [
            "Low humanization: candidates often interact with static forms or simple bots.",
            "Limited soft-skill observation compared with realistic conversational interviewing.",
            "Scalability improves, but fairness and engagement can still vary across platforms.",
            "Most systems do not combine avatar personalization, integrity checks, and resilient fallback behavior in one workflow.",
        ],
    )

    add_content_slide(
        prs,
        "Section 5",
        "Research Gaps",
        "Identified Gaps",
        [
            "Personalized interviewer avatars are still underexplored in the recruitment domain.",
            "There is a gap between scalable automation and a realistic human-like interview experience.",
            "Many systems evaluate content but do not integrate integrity monitoring into the same platform.",
            "Reliability during third-party AI outages is often overlooked in research prototypes.",
        ],
        "Gap Addressed by This Project",
        [
            "Brings avatar-based interaction, automated scoring, media capture, and ranking into one system.",
            "Adds client-side phone-use detection to strengthen interview trustworthiness.",
            "Introduces graceful degradation through local rule-based analysis when the LLM is unavailable.",
            "Shows a practical design path for scalable, engaging, and robust AI-assisted hiring.",
        ],
    )

    add_content_slide(
        prs,
        "Section 6",
        "Conclusion Slide",
        "Takeaway",
        [
            "The proposed platform addresses key limitations of traditional and existing automated interview systems.",
            "It balances scalability, objectivity, engagement, and operational resilience.",
            "The research gap is centered on humanized AI interviewing with dependable evaluation and integrity support.",
            "This presentation can be used as a concise review deck for the project report or viva discussion.",
        ],
    )

    prs.save(OUTPUT)


if __name__ == "__main__":
    build()
