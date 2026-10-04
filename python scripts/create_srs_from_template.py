"""Create a plain student-style SRS from the supplied legacy Word template."""

import os
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = r"D:\interview"
TEMPLATE_DOC = os.path.join(ROOT, "docs", "srs_template.doc")
TEMPLATE_DOCX = os.path.join(ROOT, "docs", "srs_template_converted.docx")
OUTPUT = os.path.join(ROOT, "docs", "AI_Interview_Platform_SRS_Final.docx")
DIAGRAMS = os.path.join(ROOT, "diagrams")


def convert_template():
    """Use locally installed Word to convert the legacy .doc template once."""
    if os.path.exists(TEMPLATE_DOCX):
        return
    command = (
        "$word = New-Object -ComObject Word.Application; "
        "$word.Visible = $false; "
        f"$doc = $word.Documents.Open('{TEMPLATE_DOC}'); "
        f"$doc.SaveAs2('{TEMPLATE_DOCX}', 16); "
        "$doc.Close(); $word.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", command], check=True)


def clear_template_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_font(paragraph, size=11, bold=False, italic=False):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic


def add_paragraph(doc, text="", bold=False, italic=False, align=None):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.italic = italic
    return paragraph


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.bold = True
    return paragraph


def add_bullets(doc, entries):
    for entry in entries:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.15)
        run = paragraph.add_run(f"- {entry}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def add_contents(doc):
    """Add a simple contents page appropriate for a student project report."""
    entries = [
        ("1. Introduction", 0),
        ("1.1 Purpose", 1),
        ("1.2 Document Conventions", 1),
        ("1.3 Intended Audience", 1),
        ("1.4 Project Scope", 1),
        ("1.5 References", 1),
        ("2. Overall Description", 0),
        ("3. System Features", 0),
        ("4. External Interface Requirements", 0),
        ("5. Other Nonfunctional Requirements", 0),
        ("6. Other Requirements", 0),
        ("Appendix A: Glossary", 0),
        ("Appendix B: Analysis Models", 0),
        ("B.1 Use Case Diagram", 1),
        ("B.2 Activity Diagram", 1),
        ("Appendix C: Issues List", 0),
    ]
    for text, indent in entries:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25 * indent)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def format_table(table):
    # The legacy template has its own table styles, so do not require a newer
    # Word style such as "Table Grid".
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    format_table(table)
    doc.add_paragraph()
    return table


def add_diagram(doc, filename, caption, width):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(os.path.join(DIAGRAMS, filename), width=Inches(width))
    caption_paragraph = add_paragraph(doc, caption, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    caption_paragraph.paragraph_format.space_after = Pt(10)


def build_srs():
    convert_template()
    doc = Document(TEMPLATE_DOCX)
    clear_template_body(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

    # Title page follows the supplied template fields without leaving any placeholders.
    add_paragraph(doc, "SOFTWARE REQUIREMENTS SPECIFICATION", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(16)
    add_paragraph(doc, "AI INTERVIEW PLATFORM", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(16)
    add_paragraph(doc, "Version 1.0", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "Submitted as a Student Project Report", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "Department of Computer Science", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "September 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    add_heading(doc, "Revision History", 1)
    add_table(doc, ["Name", "Date", "Reason for Changes", "Version"], [
        ("Project Team", "September 2026", "Initial SRS submission", "1.0"),
    ])

    add_heading(doc, "Table of Contents", 1)
    add_contents(doc)
    doc.add_page_break()

    add_heading(doc, "1. Introduction", 1)
    add_heading(doc, "1.1 Purpose", 2)
    add_paragraph(doc, "This Software Requirements Specification (SRS) describes Version 1.0 of the AI Interview Platform. The platform supports online interview creation, candidate registration, AI-assisted question generation, recorded interview responses, answer scoring, and integrity monitoring.")
    add_heading(doc, "1.2 Document Conventions", 2)
    add_paragraph(doc, "Requirements are written using the word 'shall' and are identified by a unique requirement ID. Use cases are identified as UC1, UC2, and so on. Functional requirements are grouped by system feature.")
    add_heading(doc, "1.3 Intended Audience and Reading Suggestions", 2)
    add_paragraph(doc, "This document is intended for the student development team, project guide, testers, and interview administrators. Readers may begin with Sections 1 and 2 for an overview, then review Section 3 for features and Appendix B for the UML diagrams.")
    add_heading(doc, "1.4 Project Scope", 2)
    add_paragraph(doc, "The AI Interview Platform reduces the manual effort of preliminary interviews. An interviewer creates a session from a job description and shares a link. A candidate registers, optionally uploads a resume, gives consent, attends an AI avatar interview, and receives evaluation through the system. The interviewer can review scores, recordings, integrity signals, and ranked results.")
    add_heading(doc, "1.5 References", 2)
    add_bullets(doc, [
        "IEEE Std 830-1998, Recommended Practice for Software Requirements Specifications.",
        "Project source code and API modules in the AI Interview Platform repository.",
        "W3C MediaDevices, MediaRecorder, and Web Speech API documentation.",
    ])

    add_heading(doc, "2. Overall Description", 1)
    add_heading(doc, "2.1 Product Perspective", 2)
    add_paragraph(doc, "The product is a browser-based web application. The frontend provides separate interviewer and candidate views. Node.js API endpoints manage sessions, candidates, question generation, response analysis, and integrity events. Data is stored through MongoDB when available, with local JSON storage as a fallback.")
    add_heading(doc, "2.2 Product Features", 2)
    add_bullets(doc, [
        "Create interview sessions and generate questions from a job description.",
        "Configure interview media and share a unique interview link.",
        "Register candidates and accept an optional resume upload.",
        "Conduct AI avatar interviews with video/audio recording and speech transcription.",
        "Score candidate answers, monitor integrity signals, and show ranked results.",
    ])
    add_heading(doc, "2.3 User Classes and Characteristics", 2)
    add_table(doc, ["User class", "Characteristics and permitted actions"], [
        ("Interviewer / HR", "Creates and configures sessions, shares links, and reviews results."),
        ("Candidate", "Registers through a shared link, uploads a resume if required, provides consent, and attends the interview."),
        ("AI Service", "Generates questions, analyzes answers, and produces scores."),
        ("Integrity Service", "Collects integrity events such as phone or gaze anomalies."),
    ])
    add_heading(doc, "2.4 Operating Environment", 2)
    add_paragraph(doc, "The system runs in a current desktop or mobile web browser with internet access. Candidate interview features require a webcam, microphone, and browser permission. The application uses a React frontend, Node.js backend services, and MongoDB or JSON data storage.")
    add_heading(doc, "2.5 Design and Implementation Constraints", 2)
    add_bullets(doc, [
        "The system depends on browser support for camera, microphone, and speech recognition APIs.",
        "AI API credentials shall remain on the server and shall not be exposed in browser code.",
        "The application shall provide basic fallback behavior when an external AI service is unavailable.",
    ])
    add_heading(doc, "2.6 User Documentation", 2)
    add_paragraph(doc, "The project will provide a short interviewer guide, a candidate instruction screen before the interview, and clear permission messages for camera and microphone access.")
    add_heading(doc, "2.7 Assumptions and Dependencies", 2)
    add_paragraph(doc, "The platform assumes candidates have a stable network connection and a supported browser. Question generation and advanced scoring depend on the configured AI service. Speech transcription depends on the availability of the browser speech service.")

    add_heading(doc, "3. System Features", 1)
    add_heading(doc, "3.1 Interview Session Management", 2)
    add_paragraph(doc, "Description and Priority: This high-priority feature allows an interviewer to create a session, enter a job description, generate questions, configure media, and share an interview link.")
    add_paragraph(doc, "Stimulus/Response Sequence: The interviewer enters the session data. The system validates the input, generates or accepts questions, saves the session, and returns a shareable link.")
    add_table(doc, ["ID", "Functional Requirement"], [
        ("FR-01", "The system shall allow an interviewer to create an interview session with a title and job description."),
        ("FR-02", "The system shall generate interview questions from the job description."),
        ("FR-03", "The system shall allow an interviewer to configure avatar and media settings."),
        ("FR-04", "The system shall generate a unique link for a saved interview session."),
    ])
    add_heading(doc, "3.2 Candidate Registration and Interview", 2)
    add_paragraph(doc, "Description and Priority: This high-priority feature supports candidate onboarding and the interview process.")
    add_paragraph(doc, "Stimulus/Response Sequence: The candidate opens the shared link, enters profile details, optionally uploads a resume, provides integrity consent, and starts the interview.")
    add_table(doc, ["ID", "Functional Requirement"], [
        ("FR-05", "The system shall register a candidate against the selected interview session."),
        ("FR-06", "The system shall allow an optional resume upload during candidate registration."),
        ("FR-07", "The system shall record candidate video and audio after consent is granted."),
        ("FR-08", "The system shall transcribe candidate speech during the interview when browser support is available."),
    ])
    add_heading(doc, "3.3 Evaluation and Integrity Monitoring", 2)
    add_paragraph(doc, "Description and Priority: This high-priority feature evaluates interview answers and provides the interviewer with useful result information.")
    add_paragraph(doc, "Stimulus/Response Sequence: After responses are captured, the system analyzes transcripts and expected answer criteria, applies integrity data, stores the result, and updates the leaderboard.")
    add_table(doc, ["ID", "Functional Requirement"], [
        ("FR-09", "The system shall score candidate answers using generated transcripts and evaluation criteria."),
        ("FR-10", "The system shall record integrity signals during an interview."),
        ("FR-11", "The system shall calculate and display candidate results and ranking information."),
        ("FR-12", "The system shall allow the interviewer to review candidate recordings, scores, and integrity information."),
    ])

    add_heading(doc, "4. External Interface Requirements", 1)
    add_heading(doc, "4.1 User Interfaces", 2)
    add_paragraph(doc, "The interviewer interface shall provide session creation, question configuration, candidate lists, and result review. The candidate interface shall provide registration, consent, camera and microphone permission prompts, question display, and response recording controls.")
    add_heading(doc, "4.2 Hardware Interfaces", 2)
    add_paragraph(doc, "The candidate device shall interface with a webcam and microphone through the browser MediaDevices API. The system shall show a clear message if either device is unavailable or permission is denied.")
    add_heading(doc, "4.3 Software Interfaces", 2)
    add_paragraph(doc, "The frontend communicates with Node.js REST APIs. The backend interfaces with the AI question and scoring service, storage service, and browser speech APIs. Data exchange uses JSON.")
    add_heading(doc, "4.4 Communications Interfaces", 2)
    add_paragraph(doc, "The application shall use HTTPS for deployed web traffic. API requests and responses shall use JSON over HTTP. Media permissions and browser recording are managed locally by the browser.")

    add_heading(doc, "5. Other Nonfunctional Requirements", 1)
    add_heading(doc, "5.1 Performance Requirements", 2)
    add_paragraph(doc, "The system should display normal application screens within three seconds on a stable connection. Question generation and result analysis should show a loading state while the external AI service is processing a request.")
    add_heading(doc, "5.2 Safety Requirements", 2)
    add_paragraph(doc, "The system shall request consent before recording candidate media. The candidate shall be informed when camera or microphone access is active.")
    add_heading(doc, "5.3 Security Requirements", 2)
    add_paragraph(doc, "The system shall protect interview links and candidate data. API keys shall be stored on the server. Candidate media and personal data shall not be exposed to unauthorized users.")
    add_heading(doc, "5.4 Software Quality Attributes", 2)
    add_bullets(doc, [
        "Usability: the candidate interview flow shall be understandable without special training.",
        "Reliability: the application shall handle unavailable AI services with a fallback response or clear error message.",
        "Maintainability: frontend pages, service modules, and API handlers shall remain separated by responsibility.",
        "Portability: the application shall run in modern browsers without installing separate client software.",
    ])

    add_heading(doc, "6. Other Requirements", 1)
    add_paragraph(doc, "Session data shall include the interview title, job description, question set, configuration, candidate records, scores, and integrity events. Candidate records shall be associated with the session link used during registration. The system shall preserve enough result information for the interviewer to compare candidates.")

    add_heading(doc, "Appendix A: Glossary", 1)
    add_table(doc, ["Term", "Meaning"], [
        ("AI Avatar", "A virtual interviewer that presents questions to the candidate."),
        ("Integrity Signal", "An event used to indicate a possible issue during an interview."),
        ("STT", "Speech-to-Text conversion of spoken answers into text."),
        ("TTS", "Text-to-Speech output used to present interview questions."),
        ("UML", "Unified Modeling Language used for system diagrams."),
    ])

    add_heading(doc, "Appendix B: Analysis Models", 1)
    add_heading(doc, "B.1 Use Case Diagram", 2)
    add_paragraph(doc, "The following diagram shows the actors, use cases, and main include/extend relationships of the AI Interview Platform.")
    add_diagram(doc, "diagram_usecase_perfect.png", "Figure B.1: AI Interview Platform Use Case Diagram", 6.8)
    add_heading(doc, "B.2 Activity Diagram", 2)
    add_paragraph(doc, "The following activity diagram shows the main flow from session creation to final results, including the optional resume path and parallel capture activities.")
    add_diagram(doc, "diagram_activity_perfect.png", "Figure B.2: AI Interview Platform Activity Diagram", 5.8)

    add_heading(doc, "Appendix C: Issues List", 1)
    add_paragraph(doc, "No open SRS issues are recorded for Version 1.0. Future work may include additional role-based access control, automated notifications, and expanded integrity analysis.")

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_srs()
